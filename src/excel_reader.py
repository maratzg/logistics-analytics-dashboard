from __future__ import annotations

import posixpath
import zipfile
from dataclasses import dataclass, field
from xml.etree import ElementTree
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from .models import SheetInfo, TableInfo, WorkbookStructure
from .schema import EXPECTED_HISTORY_COLUMNS, EXPECTED_MASTER_COLUMNS, EXPECTED_WEEKLY_COLUMNS
from .validators import ValidationIssue, workbook_schema_issues


WEEKLY_SHEET = "WEEKLY"
LOG_HISTORY_SHEET = "LOG_HISTORY"
MASTER_DATA_SHEET = "MASTER_DATA"
TEMPLATE_SHEET = "TEMPLATE"
SUPPORTED_WORKBOOK_EXTENSIONS = {".xlsm", ".xlsx"}


@dataclass(frozen=True, slots=True)
class WorkbookReadSnapshot:
    structure: WorkbookStructure
    weekly_df: pd.DataFrame
    history_df: pd.DataFrame
    master_data_df: pd.DataFrame
    read_issues: list[ValidationIssue] = field(default_factory=list)
    schema_issues: list[ValidationIssue] = field(default_factory=list)
    weekly_headers: list[str] = field(default_factory=list)
    history_headers: list[str] = field(default_factory=list)
    template_headers: list[str] = field(default_factory=list)
    master_headers: list[str] = field(default_factory=list)


class ExcelWorkbookReader:
    """Read-only data access for the macro-enabled logistics workbook."""

    def __init__(self, workbook_path: str | Path) -> None:
        self.workbook_path = Path(workbook_path).expanduser()

    def inspect_workbook(self) -> WorkbookStructure:
        self._ensure_workbook_exists()
        table_info = _read_table_info_by_sheet(self.workbook_path)
        workbook = load_workbook(
            self.workbook_path,
            read_only=True,
            data_only=False,
            keep_vba=False,
        )
        try:
            return _structure_from_workbook(workbook, table_info)
        finally:
            workbook.close()

    def read_snapshot(self) -> WorkbookReadSnapshot:
        """Read structure and all supported datasets in one workbook session."""

        self._ensure_workbook_exists()
        table_info = _read_table_info_by_sheet(self.workbook_path)
        workbook = load_workbook(
            self.workbook_path,
            read_only=True,
            data_only=True,
            keep_vba=False,
        )
        try:
            structure = _structure_from_workbook(workbook, table_info)
            read_issues: list[ValidationIssue] = []

            if WEEKLY_SHEET in workbook.sheetnames:
                weekly_ws = workbook[WEEKLY_SHEET]
                weekly_rows, issues = _read_weekly_rows(weekly_ws, table_info.get(WEEKLY_SHEET, []))
                read_issues.extend(issues)
                weekly_df = pd.DataFrame(weekly_rows, columns=_weekly_columns(weekly_rows))
                weekly_headers = _recognized_headers(weekly_ws, table_info.get(WEEKLY_SHEET, []), EXPECTED_WEEKLY_COLUMNS, 10)
            else:
                weekly_df = _empty_weekly_frame()
                weekly_headers = []
                read_issues.append(ValidationIssue("ERROR", "WEEKLY sheet is missing."))

            if LOG_HISTORY_SHEET in workbook.sheetnames:
                history_ws = workbook[LOG_HISTORY_SHEET]
                history_rows, issues = _read_history_rows(history_ws, table_info.get(LOG_HISTORY_SHEET, []))
                read_issues.extend(issues)
                history_df = pd.DataFrame(history_rows, columns=EXPECTED_HISTORY_COLUMNS if not history_rows else list(history_rows[0].keys()))
                history_headers = _recognized_headers(history_ws, table_info.get(LOG_HISTORY_SHEET, []), EXPECTED_HISTORY_COLUMNS, 8)
            else:
                history_df = pd.DataFrame(columns=EXPECTED_HISTORY_COLUMNS)
                history_headers = []
                read_issues.append(ValidationIssue("ERROR", "LOG_HISTORY sheet is missing."))

            template_headers: list[str] = []
            if TEMPLATE_SHEET in workbook.sheetnames:
                template_ws = workbook[TEMPLATE_SHEET]
                template_headers = _recognized_headers(template_ws, table_info.get(TEMPLATE_SHEET, []), EXPECTED_WEEKLY_COLUMNS, 10)

            master_headers: list[str] = []
            master_df = pd.DataFrame(columns=EXPECTED_MASTER_COLUMNS)
            if MASTER_DATA_SHEET in workbook.sheetnames:
                master_ws = workbook[MASTER_DATA_SHEET]
                master_headers = _recognized_headers(master_ws, table_info.get(MASTER_DATA_SHEET, []), EXPECTED_MASTER_COLUMNS, 10)
                master_rows = _read_master_rows(master_ws, table_info.get(MASTER_DATA_SHEET, []))
                master_df = pd.DataFrame(master_rows, columns=_weekly_columns(master_rows))

            schema_issues = workbook_schema_issues(
                weekly_headers,
                history_headers,
                template_headers=template_headers,
                master_headers=master_headers,
            )
            return WorkbookReadSnapshot(
                structure=structure,
                weekly_df=weekly_df,
                history_df=history_df,
                master_data_df=master_df,
                read_issues=read_issues,
                schema_issues=schema_issues,
                weekly_headers=weekly_headers,
                history_headers=history_headers,
                template_headers=template_headers,
                master_headers=master_headers,
            )
        finally:
            workbook.close()

    def verify_required_sheets(self, required_sheets: list[str]) -> list[ValidationIssue]:
        structure = self.inspect_workbook()
        available = {sheet.upper() for sheet in structure.sheet_names}
        issues = []
        for sheet in required_sheets:
            if sheet.upper() not in available:
                issues.append(ValidationIssue("ERROR", f"Required sheet is missing: {sheet}"))
        return issues

    def read_weekly(self) -> tuple[pd.DataFrame, list[ValidationIssue]]:
        snapshot = self.read_snapshot()
        issues = [issue for issue in snapshot.read_issues if "WEEKLY" in issue.message or issue.location]
        return snapshot.weekly_df, issues

    def read_log_history(self) -> tuple[pd.DataFrame, list[ValidationIssue]]:
        snapshot = self.read_snapshot()
        issues = [issue for issue in snapshot.read_issues if "LOG_HISTORY" in issue.message]
        return snapshot.history_df, issues

    def _ensure_workbook_exists(self) -> None:
        if not self.workbook_path.exists():
            raise FileNotFoundError(f"Workbook not found: {self.workbook_path}")
        if self.workbook_path.suffix.lower() not in SUPPORTED_WORKBOOK_EXTENSIONS:
            raise ValueError(f"Expected an Excel workbook ({', '.join(sorted(SUPPORTED_WORKBOOK_EXTENSIONS))}), got: {self.workbook_path}")


def _read_weekly_rows(worksheet: Worksheet, tables: list[TableInfo]) -> tuple[list[dict[str, Any]], list[ValidationIssue]]:
    issues: list[ValidationIssue] = []
    table_rows = _read_table_like_rows(worksheet, tables, expected_headers=EXPECTED_WEEKLY_COLUMNS)
    if table_rows:
        return table_rows, issues

    header_rows = _find_header_rows(worksheet, EXPECTED_WEEKLY_COLUMNS, minimum_matches=10)
    if not header_rows:
        non_empty_cells = [
            cell.coordinate
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value not in (None, "")
        ]
        if non_empty_cells:
            issues.append(
                ValidationIssue(
                    "WARNING",
                    "WEEKLY has cell content but no recognized weekly table/header block. No operational rows were loaded.",
                    f"First populated cells: {', '.join(non_empty_cells[:5])}",
                )
            )
        else:
            issues.append(ValidationIssue("WARNING", "WEEKLY is empty. No operational rows were loaded."))
        return _empty_weekly_records(), issues

    rows: list[dict[str, Any]] = []
    header_rows_with_end = header_rows + [worksheet.max_row + 1]
    for header_index, header_row in enumerate(header_rows):
        next_header_row = header_rows_with_end[header_index + 1]
        header_map = _header_map_for_row(worksheet, header_row, EXPECTED_WEEKLY_COLUMNS)
        for row_number in range(header_row + 1, next_header_row):
            row_dict = {column: worksheet.cell(row=row_number, column=header_map[column]).value for column in header_map}
            if not any(value not in (None, "") for value in row_dict.values()):
                continue
            row_dict["_source_excel_row"] = row_number
            row_dict["_source_header_row"] = header_row
            rows.append(row_dict)
    return rows, issues


def _read_history_rows(worksheet: Worksheet, tables: list[TableInfo]) -> tuple[list[dict[str, Any]], list[ValidationIssue]]:
    issues: list[ValidationIssue] = []
    history_table = next((table for table in tables if table.name == "tblHistory"), None)
    if history_table is not None:
        rows = _read_table_rows(worksheet, history_table.ref)
        return rows, issues

    header_rows = _find_header_rows(worksheet, EXPECTED_HISTORY_COLUMNS, minimum_matches=8)
    if not header_rows:
        issues.append(ValidationIssue("ERROR", "LOG_HISTORY has no tblHistory table and no recognizable history header row."))
        return [], issues
    rows = _read_rows_from_header(worksheet, header_rows[0], EXPECTED_HISTORY_COLUMNS)
    return rows, issues


def _read_master_rows(worksheet: Worksheet, tables: list[TableInfo]) -> list[dict[str, Any]]:
    master_table = next((table for table in tables if table.name == "tblMasterData"), None)
    if master_table is not None:
        return _read_table_rows(worksheet, master_table.ref)
    header_rows = _find_header_rows(worksheet, EXPECTED_MASTER_COLUMNS, minimum_matches=10)
    if not header_rows:
        return []
    return _read_rows_from_header(worksheet, header_rows[0], EXPECTED_MASTER_COLUMNS)


def _read_table_like_rows(worksheet: Worksheet, tables: list[TableInfo], expected_headers: list[str]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for table in tables:
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        headers = table.headers or [_clean_header(worksheet.cell(row=min_row, column=column).value) for column in range(min_col, max_col + 1)]
        matches = len(set(_normalize_header(header) for header in headers) & set(_normalize_header(h) for h in expected_headers))
        if matches < 10:
            continue
        for row_number in range(min_row + 1, max_row + 1):
            row_dict = {
                headers[column_offset]: worksheet.cell(row=row_number, column=min_col + column_offset).value
                for column_offset in range(len(headers))
                if headers[column_offset]
            }
            if not any(value not in (None, "") for value in row_dict.values()):
                continue
            row_dict["_source_excel_row"] = row_number
            row_dict["_source_table"] = table.name
            rows.append(row_dict)
    return rows


def _read_table_rows(worksheet: Worksheet, table_ref: str) -> list[dict[str, Any]]:
    min_col, min_row, max_col, max_row = range_boundaries(table_ref)
    headers = [_clean_header(worksheet.cell(row=min_row, column=column).value) for column in range(min_col, max_col + 1)]
    rows: list[dict[str, Any]] = []
    for row_number in range(min_row + 1, max_row + 1):
        row_dict = {
            headers[column_offset]: worksheet.cell(row=row_number, column=min_col + column_offset).value
            for column_offset in range(len(headers))
            if headers[column_offset]
        }
        if not any(value not in (None, "") for value in row_dict.values()):
            continue
        row_dict["_source_excel_row"] = row_number
        rows.append(row_dict)
    return rows


def _read_rows_from_header(worksheet: Worksheet, header_row: int, expected_headers: list[str]) -> list[dict[str, Any]]:
    header_map = _header_map_for_row(worksheet, header_row, expected_headers)
    rows = []
    for row_number in range(header_row + 1, worksheet.max_row + 1):
        row_dict = {column: worksheet.cell(row=row_number, column=column_index).value for column, column_index in header_map.items()}
        if not any(value not in (None, "") for value in row_dict.values()):
            continue
        row_dict["_source_excel_row"] = row_number
        rows.append(row_dict)
    return rows


def _find_header_rows(worksheet: Worksheet, expected_headers: list[str], minimum_matches: int) -> list[int]:
    normalized_expected = {_normalize_header(header) for header in expected_headers}
    matches: list[int] = []
    for row in worksheet.iter_rows():
        values = {_normalize_header(cell.value) for cell in row if cell.value not in (None, "")}
        if len(values & normalized_expected) >= minimum_matches:
            matches.append(row[0].row)
    return matches


def _header_map_for_row(worksheet: Worksheet, row_number: int, expected_headers: list[str]) -> dict[str, int]:
    normalized_expected = {_normalize_header(header): header for header in expected_headers}
    mapping = {}
    for cell in worksheet[row_number]:
        normalized = _normalize_header(cell.value)
        if normalized in normalized_expected:
            mapping[normalized_expected[normalized]] = cell.column
    return mapping


def _empty_weekly_frame() -> pd.DataFrame:
    return pd.DataFrame(columns=EXPECTED_WEEKLY_COLUMNS + ["_source_excel_row", "_source_header_row", "_source_table"])


def _empty_weekly_records() -> list[dict[str, Any]]:
    return []


def _weekly_columns(rows: list[dict[str, Any]]) -> list[str]:
    if not rows:
        return EXPECTED_WEEKLY_COLUMNS + ["_source_excel_row", "_source_header_row", "_source_table"]
    columns: list[str] = []
    for row in rows:
        for column in row:
            if column not in columns:
                columns.append(column)
    return columns


def _structure_from_workbook(workbook: Any, table_info: dict[str, list[TableInfo]]) -> WorkbookStructure:
    sheets = [
        SheetInfo(
            name=worksheet.title,
            state=worksheet.sheet_state,
            max_row=worksheet.max_row,
            max_column=worksheet.max_column,
            tables=table_info.get(worksheet.title, []),
        )
        for worksheet in workbook.worksheets
    ]
    return WorkbookStructure(list(workbook.sheetnames), sheets)


def _recognized_headers(
    worksheet: Worksheet,
    tables: list[TableInfo],
    expected_headers: list[str],
    minimum_matches: int,
) -> list[str]:
    normalized_expected = {_normalize_header(header) for header in expected_headers}
    for table in tables:
        headers = table.headers
        matches = len({_normalize_header(header) for header in headers} & normalized_expected)
        if matches >= minimum_matches:
            return headers
    header_rows = _find_header_rows(worksheet, expected_headers, minimum_matches)
    if not header_rows:
        return []
    mapping = _header_map_for_row(worksheet, header_rows[0], expected_headers)
    return list(mapping.keys())


def _read_table_info_by_sheet(workbook_path: Path) -> dict[str, list[TableInfo]]:
    """Read Excel table metadata directly from the XLSX/XLSM package.

    OpenPyXL's read-only worksheets intentionally do not expose table objects.
    Reading the table XML keeps this project out of workbook write mode while
    preserving support for existing structured tables such as tblHistory.
    """

    try:
        with zipfile.ZipFile(workbook_path) as archive:
            sheet_targets = _sheet_targets_by_name(archive)
            table_info: dict[str, list[TableInfo]] = {}
            for sheet_name, sheet_path in sheet_targets.items():
                table_info[sheet_name] = _tables_for_sheet(archive, sheet_path)
            return table_info
    except Exception:
        return {}


def _sheet_targets_by_name(archive: zipfile.ZipFile) -> dict[str, str]:
    main_ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    office_rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

    workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    rel_root = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    workbook_rels = {
        relationship.attrib["Id"]: _normalize_archive_path("xl", relationship.attrib["Target"])
        for relationship in rel_root
        if "Id" in relationship.attrib and "Target" in relationship.attrib
    }

    targets: dict[str, str] = {}
    sheets = workbook_root.find(f"{main_ns}sheets")
    if sheets is None:
        return targets
    for sheet in sheets.findall(f"{main_ns}sheet"):
        rel_id = sheet.attrib.get(f"{office_rel_ns}id")
        name = sheet.attrib.get("name")
        if rel_id and name and rel_id in workbook_rels:
            targets[name] = workbook_rels[rel_id]
    return targets


def _tables_for_sheet(archive: zipfile.ZipFile, sheet_path: str) -> list[TableInfo]:
    main_ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    office_rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    package_rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"

    try:
        sheet_root = ElementTree.fromstring(archive.read(sheet_path))
    except KeyError:
        return []

    table_part_ids = [
        node.attrib.get(f"{office_rel_ns}id")
        for node in sheet_root.findall(f".//{main_ns}tablePart")
        if node.attrib.get(f"{office_rel_ns}id")
    ]
    if not table_part_ids:
        return []

    rel_path = _relationship_path(sheet_path)
    try:
        rel_root = ElementTree.fromstring(archive.read(rel_path))
    except KeyError:
        return []

    table_targets = {
        relationship.attrib.get("Id"): _normalize_archive_path(posixpath.dirname(sheet_path), relationship.attrib.get("Target", ""))
        for relationship in rel_root.findall(f"{package_rel_ns}Relationship")
    }

    tables: list[TableInfo] = []
    for rel_id in table_part_ids:
        target = table_targets.get(rel_id)
        if not target:
            continue
        table = _read_table_info(archive, target)
        if table is not None:
            tables.append(table)
    return tables


def _read_table_info(archive: zipfile.ZipFile, table_path: str) -> TableInfo | None:
    main_ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

    try:
        root = ElementTree.fromstring(archive.read(table_path))
    except KeyError:
        return None

    name = root.attrib.get("displayName") or root.attrib.get("name")
    table_ref = root.attrib.get("ref")
    if not name or not table_ref:
        return None

    headers = [
        _clean_header(column.attrib.get("name"))
        for column in root.findall(f".//{main_ns}tableColumn")
        if _clean_header(column.attrib.get("name"))
    ]
    return TableInfo(name=name, ref=table_ref, headers=headers)


def _relationship_path(source_path: str) -> str:
    directory = posixpath.dirname(source_path)
    filename = posixpath.basename(source_path)
    return posixpath.join(directory, "_rels", f"{filename}.rels")


def _normalize_archive_path(base_directory: str, target: str) -> str:
    if target.startswith("/"):
        return posixpath.normpath(target.lstrip("/"))
    return posixpath.normpath(posixpath.join(base_directory, target))


def _clean_header(value: object) -> str:
    return " ".join(str(value or "").replace("\n", " ").split()).strip()


def _normalize_header(value: object) -> str:
    return _clean_header(value).lower()
