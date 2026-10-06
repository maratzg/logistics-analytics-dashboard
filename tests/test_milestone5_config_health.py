from __future__ import annotations

import json
import logging
import tempfile
import unittest
from logging.handlers import RotatingFileHandler
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from src.app_state import ApplicationState
from src.config import AppConfig, load_config, save_workbook_path
from src.excel_reader import EXPECTED_WEEKLY_COLUMNS
from src.logging_setup import configure_logging
from src.validators import EXPECTED_HISTORY_COLUMNS
from src.version import APP_VERSION
from src.workbook_discovery import discover_workbooks
from src.workbook_health import WorkbookHealthStatus, validate_workbook


def write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def make_workbook(path: Path, *, sheets: list[str] | None = None, vessel: str = "VESSEL A", include_history_table: bool = False) -> Path:
    requested = sheets or ["WEEKLY", "LOG_HISTORY"]
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in requested:
        worksheet = workbook.create_sheet(sheet_name)
        if sheet_name == "WEEKLY":
            worksheet.append(EXPECTED_WEEKLY_COLUMNS)
            row = {
                "Week": 35,
                "Vessel": vessel,
                "Voyage": "V001",
                "ETA": "2026-09-05",
                "ETD": "2026-09-06",
                "Cut-Off": "2026-09-04",
                "ETA T/S": "2026-09-10",
                "Booking number": "BK001",
                "CNTR AMT": 1,
                "SIZE": 2,
                "GWT": 10,
                "CUSTOMER": "Customer",
                "POD": "RIX",
                "RowStatus": "",
                "CancelReason": "",
                "Summary": 1,
                "TEU": 1,
                "TS": 10,
                "BlockID": "2026-W35",
                "RowID": "2026-W35|R01",
            }
            worksheet.append([row.get(column, "") for column in EXPECTED_WEEKLY_COLUMNS])
        elif sheet_name == "LOG_HISTORY":
            worksheet.append(EXPECTED_HISTORY_COLUMNS)
            if include_history_table:
                worksheet.append([""] * len(EXPECTED_HISTORY_COLUMNS))
                table = Table(displayName="tblHistory", ref=f"A1:P{worksheet.max_row}")
                worksheet.add_table(table)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


class ConfigPersistenceTests(unittest.TestCase):
    def test_missing_user_config_loads_project_defaults(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            project_config = Path(temp) / "project" / "config.json"
            user_config = Path(temp) / "user" / "config.json"
            write_json(project_config, {"required_sheets": ["WEEKLY", "LOG_HISTORY"], "max_issue_details": 12})

            config = load_config(project_config, user_config_path=user_config)

            self.assertIsNone(config.workbook_path)
            self.assertEqual(config.max_issue_details, 12)
            self.assertEqual(config.required_sheets, ["WEEKLY", "LOG_HISTORY"])

    def test_config_creation_and_read_with_spaces_and_unicode_path(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            project_config = Path(temp) / "config.json"
            user_config = Path(temp) / "Anna Bērziņa" / "config.json"
            workbook_path = Path(temp) / "OneDrive - Example Company SIA" / "Latvia Documents" / "logistics_operations_demo.xlsm"
            make_workbook(workbook_path)
            write_json(project_config, {})

            saved = save_workbook_path(workbook_path, user_config_path=user_config, existing_config=load_config(project_config, user_config_path=user_config))
            loaded = load_config(project_config, user_config_path=user_config)

            self.assertTrue(user_config.exists())
            self.assertEqual(saved.workbook_path, workbook_path)
            self.assertEqual(loaded.workbook_path, workbook_path)

    def test_invalid_config_json_is_reported_without_crashing(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            project_config = Path(temp) / "config.json"
            user_config = Path(temp) / "config-user.json"
            write_json(project_config, {})
            user_config.write_text("{ not valid json", encoding="utf-8")

            config = load_config(project_config, user_config_path=user_config)

            self.assertIsNone(config.workbook_path)
            self.assertTrue(config.config_errors)

    def test_user_config_overrides_project_default_workbook_path(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            project_workbook = make_workbook(Path(temp) / "project workbook.xlsm")
            user_workbook = make_workbook(Path(temp) / "user workbook.xlsm")
            project_config = Path(temp) / "project" / "config.json"
            user_config = Path(temp) / "user" / "config.json"
            write_json(project_config, {"workbook_path": str(project_workbook)})
            write_json(user_config, {"workbook_path": str(user_workbook)})

            config = load_config(project_config, user_config_path=user_config)

            self.assertEqual(config.workbook_path, user_workbook)

    def test_windows_style_unicode_path_round_trips_in_user_config(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            project_config = Path(temp) / "config.json"
            user_config = Path(temp) / "config-user.json"
            write_json(project_config, {})
            windows_path = r"C:\Users\Anna Bērziņa\OneDrive - Example Company SIA\Latvia Documents\logistics_operations_demo.xlsm"

            saved = save_workbook_path(windows_path, user_config_path=user_config, existing_config=load_config(project_config, user_config_path=user_config))

            self.assertEqual(str(saved.workbook_path), windows_path)
            self.assertIn("Bērziņa", user_config.read_text(encoding="utf-8"))

    def test_invalid_workbook_selection_does_not_replace_good_configuration(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            good_workbook = make_workbook(Path(temp) / "good.xlsm")
            invalid_file = Path(temp) / "not workbook.txt"
            invalid_file.write_text("nope", encoding="utf-8")
            project_config = Path(temp) / "config.json"
            user_config = Path(temp) / "user-config.json"
            write_json(project_config, {})
            config = save_workbook_path(good_workbook, user_config_path=user_config, existing_config=load_config(project_config, user_config_path=user_config))
            state = ApplicationState(config)

            result = state.change_workbook(invalid_file)
            reloaded = load_config(project_config, user_config_path=user_config)

            self.assertFalse(result.success)
            self.assertEqual(state.config.workbook_path, good_workbook)
            self.assertEqual(reloaded.workbook_path, good_workbook)


class WorkbookHealthTests(unittest.TestCase):
    def test_valid_workbook_structure_with_history_table(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            workbook = make_workbook(Path(temp) / "valid.xlsm", include_history_table=True)

            health = validate_workbook(workbook, optional_sheets=[])

            self.assertEqual(health.status, WorkbookHealthStatus.READY)
            self.assertIn("WEEKLY", health.detected_sheets)
            history_sheet = next(sheet for sheet in health.structure.sheets if sheet.name == "LOG_HISTORY")
            self.assertEqual(history_sheet.tables[0].name, "tblHistory")

    def test_missing_weekly_is_invalid(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            workbook = make_workbook(Path(temp) / "missing_weekly.xlsm", sheets=["LOG_HISTORY"])

            health = validate_workbook(workbook)

            self.assertEqual(health.status, WorkbookHealthStatus.INVALID_WORKBOOK)
            self.assertTrue(any("WEEKLY" in error for error in health.errors))

    def test_missing_log_history_is_invalid(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            workbook = make_workbook(Path(temp) / "missing_history.xlsm", sheets=["WEEKLY"])

            health = validate_workbook(workbook)

            self.assertEqual(health.status, WorkbookHealthStatus.INVALID_WORKBOOK)
            self.assertTrue(any("LOG_HISTORY" in error for error in health.errors))

    def test_unsupported_extension_is_invalid(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            file_path = Path(temp) / "workbook.csv"
            file_path.write_text("Week,Vessel", encoding="utf-8")

            health = validate_workbook(file_path)

            self.assertEqual(health.status, WorkbookHealthStatus.INVALID_WORKBOOK)
            self.assertTrue(any("Unsupported" in error for error in health.errors))

    def test_missing_file_reports_file_not_found(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            health = validate_workbook(Path(temp) / "missing.xlsm")

            self.assertEqual(health.status, WorkbookHealthStatus.FILE_NOT_FOUND)

    def test_corrupted_workbook_reports_load_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            workbook = Path(temp) / "corrupt.xlsm"
            workbook.write_bytes(b"not a workbook")

            health = validate_workbook(workbook)

            self.assertEqual(health.status, WorkbookHealthStatus.LOAD_ERROR)

    def test_optional_sheets_absent_are_warnings_not_errors(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            workbook = make_workbook(Path(temp) / "valid with warnings.xlsm")

            health = validate_workbook(workbook, optional_sheets=["README", "REPORT"])

            self.assertEqual(health.status, WorkbookHealthStatus.READY_WITH_WARNINGS)
            self.assertEqual(len(health.warnings), 2)

    def test_path_with_spaces_and_unicode_characters_validates(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            workbook = make_workbook(Path(temp) / "Anna Bērziņa" / "OneDrive - Example Company SIA" / "logistics_operations_demo.xlsm")

            health = validate_workbook(workbook)

            self.assertEqual(health.status, WorkbookHealthStatus.READY)


class ApplicationStateMilestone5Tests(unittest.TestCase):
    def test_not_configured_startup_state(self) -> None:
        state = ApplicationState(AppConfig(workbook_path=None))

        result = state.refresh()
        model = state.dashboard_model()

        self.assertFalse(result.success)
        self.assertEqual(state.health.status, WorkbookHealthStatus.NOT_CONFIGURED)
        self.assertTrue(model.setup_required)
        self.assertIn("not configured", model.setup_title.casefold())

    def test_valid_startup_loads_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            workbook = make_workbook(Path(temp) / "valid.xlsm", vessel="VESSEL READY")
            state = ApplicationState(AppConfig(workbook_path=workbook, optional_sheets=[]))

            result = state.refresh()

            self.assertTrue(result.success)
            self.assertEqual(state.health.status, WorkbookHealthStatus.READY)
            self.assertEqual(state.data.weekly_df.iloc[0]["Vessel"], "VESSEL READY")

    def test_missing_workbook_startup_state(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            state = ApplicationState(AppConfig(workbook_path=Path(temp) / "missing.xlsm"))

            result = state.refresh()

            self.assertFalse(result.success)
            self.assertEqual(state.health.status, WorkbookHealthStatus.FILE_NOT_FOUND)
            self.assertIsNone(state.data)

    def test_invalid_workbook_startup_state(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            workbook = make_workbook(Path(temp) / "invalid.xlsm", sheets=["WEEKLY"])
            state = ApplicationState(AppConfig(workbook_path=workbook))

            result = state.refresh()

            self.assertFalse(result.success)
            self.assertEqual(state.health.status, WorkbookHealthStatus.INVALID_WORKBOOK)
            self.assertIsNone(state.data)

    def test_refresh_preserves_old_data_after_configured_workbook_disappears(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            workbook = make_workbook(Path(temp) / "valid.xlsm", vessel="VESSEL OLD")
            state = ApplicationState(AppConfig(workbook_path=workbook, optional_sheets=[]))
            first = state.refresh()
            state.config = AppConfig(workbook_path=Path(temp) / "missing.xlsm", optional_sheets=[])

            second = state.refresh()

            self.assertTrue(first.success)
            self.assertFalse(second.success)
            self.assertEqual(state.data.weekly_df.iloc[0]["Vessel"], "VESSEL OLD")
            self.assertIn("previously loaded", second.message)

    def test_changing_workbook_successfully_reloads_state(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            first_workbook = make_workbook(Path(temp) / "first.xlsm", vessel="VESSEL FIRST")
            second_workbook = make_workbook(Path(temp) / "second.xlsm", vessel="VESSEL SECOND")
            config = AppConfig(workbook_path=first_workbook, optional_sheets=[], user_config_path=Path(temp) / "user" / "config.json")
            state = ApplicationState(config)
            state.refresh()

            result = state.change_workbook(second_workbook)

            self.assertTrue(result.success)
            self.assertEqual(state.config.workbook_path, second_workbook)
            self.assertEqual(state.data.weekly_df.iloc[0]["Vessel"], "VESSEL SECOND")

    def test_changing_to_invalid_workbook_preserves_old_config_and_data(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            good_workbook = make_workbook(Path(temp) / "good.xlsm", vessel="VESSEL GOOD")
            bad_workbook = make_workbook(Path(temp) / "bad.xlsm", sheets=["WEEKLY"])
            state = ApplicationState(AppConfig(workbook_path=good_workbook, optional_sheets=[], user_config_path=Path(temp) / "user" / "config.json"))
            state.refresh()

            result = state.change_workbook(bad_workbook)

            self.assertFalse(result.success)
            self.assertEqual(state.config.workbook_path, good_workbook)
            self.assertEqual(state.data.weekly_df.iloc[0]["Vessel"], "VESSEL GOOD")

    def test_settings_view_model_exposes_status_and_version(self) -> None:
        state = ApplicationState(AppConfig(workbook_path=None))
        state.refresh()

        model = state.settings_model()

        self.assertEqual(model.app_version, APP_VERSION)
        self.assertEqual(model.health_status, WorkbookHealthStatus.NOT_CONFIGURED)
        self.assertIn("Not configured", model.configured_workbook_path)


class WorkbookDiscoveryTests(unittest.TestCase):
    def test_exactly_one_match(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp) / "OneDrive - Example Company SIA"
            expected = root / "Latvia Documents" / "logistics_operations_demo.xlsm"
            expected.parent.mkdir(parents=True)
            expected.touch()

            result = discover_workbooks(roots=[root], max_depth=4)

            self.assertEqual(result.candidates, [expected])

    def test_multiple_matches_require_user_selection(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            first = root / "OneDrive - A" / "logistics_operations_demo.xlsm"
            second = root / "OneDrive - B" / "Ops" / "logistics_operations_demo.xlsm"
            first.parent.mkdir(parents=True)
            second.parent.mkdir(parents=True)
            first.touch()
            second.touch()

            result = discover_workbooks(roots=[root], max_depth=4)

            self.assertEqual(set(result.candidates), {first, second})

    def test_no_matches(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)

            result = discover_workbooks(roots=[root])

            self.assertEqual(result.candidates, [])

    def test_irrelevant_similarly_named_files_are_ignored(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            similar = root / "logistics_operations_demo copy.xlsm"
            similar.touch()

            result = discover_workbooks(roots=[root])

            self.assertEqual(result.candidates, [])

    def test_directory_depth_limit(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            too_deep = root / "a" / "b" / "c" / "logistics_operations_demo.xlsm"
            too_deep.parent.mkdir(parents=True)
            too_deep.touch()

            result = discover_workbooks(roots=[root], max_depth=1)

            self.assertEqual(result.candidates, [])

    def test_onedrive_style_path_is_found(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            profile = Path(temp) / "Users" / "Anna"
            onedrive = profile / "OneDrive - Example Company SIA"
            workbook = onedrive / "Latvia Documents" / "logistics_operations_demo.xlsm"
            workbook.parent.mkdir(parents=True)
            workbook.touch()

            result = discover_workbooks(roots=[onedrive], max_depth=3)

            self.assertEqual(result.candidates, [workbook])


class LoggingSetupTests(unittest.TestCase):
    def test_logging_initializes_rotating_file_handler(self) -> None:
        root = logging.getLogger()
        for handler in list(root.handlers):
            if getattr(handler, "_history_dashboard_file_handler", False) or getattr(handler, "_history_dashboard_console_handler", False):
                root.removeHandler(handler)
                handler.close()

        with tempfile.TemporaryDirectory() as temp:
            log_path = configure_logging(log_dir=Path(temp) / "logs", console=False)
            handlers = [handler for handler in root.handlers if getattr(handler, "_history_dashboard_file_handler", False)]

            self.assertEqual(log_path.name, "app.log")
            self.assertTrue(log_path.parent.exists())
            self.assertTrue(any(isinstance(handler, RotatingFileHandler) for handler in handlers))


if __name__ == "__main__":
    unittest.main()
