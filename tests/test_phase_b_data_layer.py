from __future__ import annotations

import hashlib
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from src.advanced_analytics import cancellations_by_week
from src.analytics import calculate_history_metrics, calculate_weekly_metrics, reconstruct_field_value
from src.app_state import ApplicationState, load_workbook_data
from src.config import AppConfig
from src.data_cleaner import clean_history, clean_weekly
from src.excel_reader import ExcelWorkbookReader
from src.filters import filter_history, filter_weekly
from src.history_processing import deduplicate_voyage_events
from src.schema import (
    CALCULATED_COLUMNS,
    EXPECTED_HISTORY_COLUMNS,
    EXPECTED_WEEKLY_COLUMNS,
    EXPANDED_HISTORY_FIELDS,
)
from src.validators import cross_check_master_data, validate_weekly
from src.workbook_health import WorkbookHealthStatus, validate_workbook


REAL_V2_FIXTURE = Path(
    os.environ.get("HISTORY_DASHBOARD_V2_WORKBOOK", "demo/logistics_operations_demo.xlsm")
)


def weekly_row(**overrides: object) -> dict[str, object]:
    row: dict[str, object] = {
        "Week": 17,
        "Vessel": "VESSEL A",
        "Voyage": "V-001",
        "ETA": "2026-04-22",
        "ETD": "2026-04-23",
        "Cut-Off": "2026-04-20",
        "ETA T/S": "2026-05-01",
        "Booking number": "BK-001",
        "CNTR AMT": 2,
        "SIZE": 2,
        "GWT": 12,
        "RowStatus": "",
        "Summary": 999,
        "TEU": 999,
        "TS": 999,
        "BlockID": "opaque-block-A",
        "RowID": "shipment/α/001",
        "LoadStatus": "",
    }
    row.update(overrides)
    return row


def history_row(**overrides: object) -> dict[str, object]:
    row: dict[str, object] = {
        "Timestamp": "2026-04-10 09:30",
        "User": "employee",
        "BlockID": "opaque-block-A",
        "RowID": "shipment/α/001",
        "Week": 17,
        "Vessel": "VESSEL A",
        "Voyage": "V-001",
        "Booking number": "BK-001",
        "Comments 1": "",
        "Comments 2": "",
        "CancelReason": "",
        "Field": "ETA",
        "OldValue": "20.04.2026",
        "NewValue": "22.04.2026",
        "Year": 1900,
        "Month": 1,
    }
    row.update(overrides)
    return row


def write_v2_workbook(path: Path, *, master_row: dict[str, object] | None = None) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)

    weekly = workbook.create_sheet("WEEKLY")
    weekly.append(EXPECTED_WEEKLY_COLUMNS)
    source = weekly_row()
    weekly.append([source.get(column) for column in EXPECTED_WEEKLY_COLUMNS])
    weekly.add_table(Table(displayName="tblWeeklyTest", ref=f"A1:AS{weekly.max_row}"))

    history = workbook.create_sheet("LOG_HISTORY")
    history.append(EXPECTED_HISTORY_COLUMNS)
    event = history_row()
    history.append([event.get(column) for column in EXPECTED_HISTORY_COLUMNS])
    history.add_table(Table(displayName="tblHistory", ref=f"A1:P{history.max_row}"))

    template = workbook.create_sheet("TEMPLATE")
    template.append(EXPECTED_WEEKLY_COLUMNS)
    template.append([None] * len(EXPECTED_WEEKLY_COLUMNS))
    template.add_table(Table(displayName="tblTemplateTest", ref="A1:AS2"))

    master = workbook.create_sheet("MASTER_DATA")
    master.append(EXPECTED_WEEKLY_COLUMNS)
    derived = master_row if master_row is not None else source
    master.append([derived.get(column) for column in EXPECTED_WEEKLY_COLUMNS])
    master.add_table(Table(displayName="tblMasterData", ref="A1:AS2"))

    workbook.create_sheet("REPORT")
    workbook.create_sheet("README")
    workbook.save(path)
    workbook.close()
    return path


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class PhaseBV2SchemaTests(unittest.TestCase):
    def test_canonical_workbook_schema_contains_exactly_45_columns(self) -> None:
        self.assertEqual(len(EXPECTED_WEEKLY_COLUMNS), 45)
        self.assertEqual(EXPECTED_WEEKLY_COLUMNS[29], "LoadStatus")
        self.assertEqual(EXPECTED_WEEKLY_COLUMNS[-1], "Total GWT")

    def test_opaque_ids_are_operational_and_week_is_independent(self) -> None:
        cleaned = clean_weekly(pd.DataFrame([weekly_row(Week=8, BlockID="not-a-week", RowID="opaque::row")])).dataframe

        self.assertTrue(bool(cleaned.iloc[0]["_is_operational"]))
        self.assertEqual(cleaned.iloc[0]["Week"], 8)
        self.assertEqual(len(filter_weekly(cleaned, week=8)), 1)
        self.assertEqual(len(filter_weekly(cleaned, week=17)), 0)

    def test_eta_year_is_explicit_and_never_parsed_from_ids(self) -> None:
        cleaned = clean_weekly(
            pd.DataFrame(
                [
                    weekly_row(ETA="2024-12-31", RowID="2099-W99|R01", BlockID="year-3000"),
                    weekly_row(ETA="invalid", RowID="2025-W01|R01", BlockID="2025-W01"),
                ]
            )
        ).dataframe

        self.assertEqual(cleaned.iloc[0]["ETA Year"], 2024)
        self.assertEqual(cleaned.iloc[1]["ETA Year"], "Unknown")
        self.assertEqual(len(filter_weekly(cleaned, year=2024)), 1)
        self.assertEqual(len(filter_weekly(cleaned, year=2099)), 0)
        self.assertEqual(len(filter_weekly(cleaned, year=2025)), 0)

    def test_week_aggregation_uses_editable_week_not_block_id(self) -> None:
        cleaned = clean_weekly(
            pd.DataFrame(
                [
                    weekly_row(Week=8, BlockID="year-3000", RowID="opaque-a", RowStatus="Cancelled"),
                    weekly_row(Week=8, BlockID="week-999", RowID="opaque-b", RowStatus="Cancelled"),
                ]
            )
        ).dataframe

        result = cancellations_by_week(cleaned)
        self.assertEqual(result.to_dict("records"), [{"Week": "8", "Cancelled rows": 2}])


class PhaseBCalculationTests(unittest.TestCase):
    def test_load_status_and_all_split_metrics_are_calculated_centrally(self) -> None:
        cleaned = clean_weekly(
            pd.DataFrame(
                [
                    weekly_row(RowID="loaded-20"),
                    weekly_row(RowID="empty-40", **{"CNTR AMT": 3, "SIZE": 4, "GWT": 18, "LoadStatus": " empty "}),
                    weekly_row(RowID="cancelled", **{"CNTR AMT": 5, "SIZE": 4, "GWT": 30, "RowStatus": "Cancelled"}),
                ]
            )
        ).dataframe

        loaded = cleaned.iloc[0]
        empty = cleaned.iloc[1]
        cancelled = cleaned.iloc[2]
        self.assertEqual(loaded["LoadStatus"], "Loaded")
        self.assertEqual(loaded["Loaded Containers"], 2)
        self.assertEqual(loaded["Loaded 20ft"], 2)
        self.assertEqual(loaded["Loaded TEU"], 2)
        self.assertEqual(loaded["Loaded GWT"], 12)
        self.assertEqual(empty["LoadStatus"], "Empty")
        self.assertEqual(empty["Empty Containers"], 3)
        self.assertEqual(empty["Empty 40ft"], 3)
        self.assertEqual(empty["Empty TEU"], 6)
        self.assertEqual(empty["Empty GWT"], 18)
        for column in CALCULATED_COLUMNS:
            self.assertEqual(cancelled[column], 0, column)

        metrics = calculate_weekly_metrics(cleaned)
        self.assertEqual(metrics.active_rows, 2)
        self.assertEqual(metrics.cancelled_rows, 1)
        self.assertEqual(metrics.loaded_containers, 2)
        self.assertEqual(metrics.empty_containers, 3)
        self.assertEqual(metrics.total_containers, 5)
        self.assertEqual(metrics.total_teu, 8)
        self.assertEqual(metrics.total_gwt, 30)

    def test_formula_cache_is_retained_but_not_trusted(self) -> None:
        cleaned = clean_weekly(pd.DataFrame([weekly_row(Summary=999, TEU=999, TS=999)])).dataframe
        row = cleaned.iloc[0]

        self.assertEqual(row["_workbook_Summary"], 999)
        self.assertEqual(row["Summary"], 2)
        self.assertEqual(row["TEU"], 2)
        self.assertEqual(row["TS"], 12)
        issues = validate_weekly(cleaned)
        self.assertTrue(any("formula cache" in issue.message for issue in issues))

    def test_unknown_load_status_is_not_silently_allocated(self) -> None:
        cleaned = clean_weekly(pd.DataFrame([weekly_row(LoadStatus="Mystery")])).dataframe
        row = cleaned.iloc[0]

        self.assertEqual(row["LoadStatus"], "Unknown")
        self.assertEqual(row["Total Containers"], 2)
        self.assertEqual(row["Total TEU"], 2)
        self.assertEqual(row["Total GWT"], 12)
        self.assertEqual(row["Unclassified Containers"], 2)
        self.assertEqual(row["Unclassified TEU"], 2)
        self.assertEqual(row["Unclassified GWT"], 12)
        self.assertEqual(row["Loaded Containers"], 0)
        self.assertEqual(row["Empty Containers"], 0)
        self.assertTrue(any("Unrecognized LoadStatus" in issue.message for issue in validate_weekly(cleaned)))


class PhaseBHistoryTests(unittest.TestCase):
    def test_history_time_dimensions_come_only_from_timestamp(self) -> None:
        raw = pd.DataFrame([history_row(Timestamp="2026-10-05 08:00", Year=1999, Month=1)])
        cleaned = clean_history(raw).dataframe

        self.assertEqual(cleaned.iloc[0]["History Year"], 2026)
        self.assertEqual(cleaned.iloc[0]["History Month"], 10)
        self.assertEqual(cleaned.iloc[0]["History Quarter"], "2026 Q4")
        self.assertEqual(len(filter_history(cleaned, year=2026, month=10, quarter="2026 Q4")), 1)
        self.assertEqual(len(filter_history(cleaned, year=1999)), 0)
        self.assertEqual(len(filter_history(raw, year=2026, month=10, quarter="2026 Q4")), 1)
        self.assertEqual(len(filter_history(raw, year=1999, month=1)), 0)

    def test_expanded_history_fields_are_supported(self) -> None:
        raw = pd.DataFrame([history_row(RowID=f"row-{index}", Field=field, OldValue="A", NewValue="B") for index, field in enumerate(EXPANDED_HISTORY_FIELDS)])
        cleaned = clean_history(raw).dataframe
        metrics = calculate_history_metrics(cleaned)

        self.assertEqual(set(metrics.field_changes), set(EXPANDED_HISTORY_FIELDS))
        self.assertTrue(all(metrics.field_changes[field] == 1 for field in EXPANDED_HISTORY_FIELDS))
        self.assertEqual(len(filter_history(cleaned, field="load status")), 1)
        self.assertEqual(len(filter_history(cleaned, field="row status")), 1)

    def test_bulk_edit_is_unavailable_and_raw_value_is_retained(self) -> None:
        history = clean_history(pd.DataFrame([history_row(OldValue="[bulk edit]", NewValue="22.04.2026")])).dataframe
        weekly = clean_weekly(pd.DataFrame([weekly_row(ETA="2026-04-22")])).dataframe
        reconstruction = reconstruct_field_value(weekly, history, "ETA", row_id="shipment/α/001")

        self.assertTrue(bool(history.iloc[0]["_legacy_bulk_edit"]))
        self.assertEqual(history.iloc[0]["_raw_OldValue"], "[bulk edit]")
        self.assertTrue(pd.isna(history.iloc[0]["OldValue"]))
        self.assertEqual(reconstruction.current_source, "WEEKLY")
        self.assertEqual(reconstruction.original_raw, reconstruction.current_raw)

    def test_voyage_event_deduplication_preserves_raw_rows_and_counts_sources(self) -> None:
        raw = pd.DataFrame(
            [
                history_row(RowID="row-1", **{"Booking number": "BK-1"}),
                history_row(RowID="row-2", **{"Booking number": "BK-2"}),
                history_row(RowID="row-3", NewValue="23.04.2026"),
            ]
        )
        cleaned = clean_history(raw).dataframe
        deduplicated = deduplicate_voyage_events(cleaned)

        self.assertEqual(len(raw), 3)
        self.assertEqual(len(cleaned), 3)
        self.assertEqual(len(deduplicated), 2)
        self.assertEqual(sorted(deduplicated["_source_event_count"].tolist()), [1, 2])

    def test_incomplete_voyage_identity_is_not_deduplicated(self) -> None:
        cleaned = clean_history(pd.DataFrame([history_row(RowID="r1", BlockID=""), history_row(RowID="r2", BlockID="")])).dataframe
        self.assertEqual(len(deduplicate_voyage_events(cleaned)), 2)


class PhaseBReaderAndHealthTests(unittest.TestCase):
    def test_single_snapshot_contains_raw_current_history_and_optional_master(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = write_v2_workbook(Path(temporary) / "v2.xlsx")
            before = sha256(path)
            config = AppConfig(workbook_path=path, required_sheets=["WEEKLY", "LOG_HISTORY"], optional_sheets=["TEMPLATE", "MASTER_DATA", "REPORT", "README"])
            data = load_workbook_data(config)

            self.assertEqual(len(data.weekly_raw_df), 1)
            self.assertEqual(len(data.weekly_df), 1)
            self.assertEqual(len(data.history_raw_df), 1)
            self.assertEqual(len(data.voyage_history_df), 1)
            self.assertEqual(len(data.master_data_raw_df), 1)
            self.assertEqual(data.weekly_df.iloc[0]["Summary"], 2)
            self.assertEqual(before, sha256(path))

    def test_normal_refresh_opens_workbook_once(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = write_v2_workbook(Path(temporary) / "v2.xlsx")
            state = ApplicationState(AppConfig(workbook_path=path, required_sheets=["WEEKLY", "LOG_HISTORY"], optional_sheets=[]))
            from src import excel_reader

            original = excel_reader.load_workbook
            with patch("src.excel_reader.load_workbook", wraps=original) as mocked:
                result = state.refresh()

            self.assertTrue(result.success)
            self.assertEqual(mocked.call_count, 1)

    def test_master_cross_check_never_replaces_weekly(self) -> None:
        weekly = clean_weekly(pd.DataFrame([weekly_row(RowID="weekly-row")])).dataframe
        master = clean_weekly(pd.DataFrame([weekly_row(RowID="other-row")])).dataframe
        issues = cross_check_master_data(weekly, master)

        self.assertTrue(any("missing" in issue.message for issue in issues))
        self.assertTrue(any("not present in WEEKLY" in issue.message for issue in issues))
        self.assertEqual(weekly.iloc[0]["RowID"], "weekly-row")

    def test_v2_schema_health_is_ready(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = write_v2_workbook(Path(temporary) / "v2.xlsx")
            health = validate_workbook(path, required_sheets=["WEEKLY", "LOG_HISTORY"], optional_sheets=["TEMPLATE", "MASTER_DATA", "REPORT", "README"])
            self.assertEqual(health.status, WorkbookHealthStatus.READY)

    @unittest.skipUnless(REAL_V2_FIXTURE.exists(), "Real V2 workbook fixture is not mounted")
    def test_real_v2_workbook_is_read_only_compatible(self) -> None:
        before = sha256(REAL_V2_FIXTURE)
        snapshot = ExcelWorkbookReader(REAL_V2_FIXTURE).read_snapshot()
        health = validate_workbook(REAL_V2_FIXTURE, required_sheets=["WEEKLY", "LOG_HISTORY"], optional_sheets=["TEMPLATE", "MASTER_DATA", "REPORT", "README"])

        self.assertEqual(snapshot.template_headers, EXPECTED_WEEKLY_COLUMNS)
        self.assertEqual(snapshot.master_headers, EXPECTED_WEEKLY_COLUMNS)
        self.assertTrue(health.can_load)
        self.assertEqual(before, sha256(REAL_V2_FIXTURE))


if __name__ == "__main__":
    unittest.main()
